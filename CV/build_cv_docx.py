#!/usr/bin/env python
"""Build the Word CV(s) from the YAML SSOT — English and Korean, with ID photo.

    python CV/build_cv_docx.py                 # both languages
    python CV/build_cv_docx.py --lang en       # English only
    python CV/build_cv_docx.py --lang ko       # Korean only
    python CV/build_cv_docx.py --out FILE      # single language, explicit path

Section order mirrors the Typst CV (CV/build_cv.py → build/cv.pdf) exactly:
Education · Experience · Publications (first / co) · Research Projects ·
Conferences (intl / domestic) · Patents · Specializations · Active Collaborations ·
Research Interests · Google Scholar Metrics.

The Korean edition (`--lang ko`) drops every publication- and conference-derived
section — publications, conference presentations, and the Scholar citation
metrics — per the user's spec, and renders everything else from the *_kr fields
in data/*.yaml.

`Curriculum Vitae Chihun Lee (2025.11).docx` is the pre-2026 baseline and is
never touched; output goes to new dated files beside it.
"""

from __future__ import annotations

import argparse
import io
from pathlib import Path

import yaml
from docx import Document
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor

ROOT = Path(__file__).resolve().parent.parent
DATA = ROOT / "data"

BODY_PT = 10
NAME_PT = 17
HEAD_PT = 11
KO_FONT = "맑은 고딕"
STAMP = "2026.09"


def load(name: str) -> dict:
    with open(DATA / name, encoding="utf-8") as f:
        return yaml.safe_load(f)


def pick(d: dict, key: str, ko: bool, default=""):
    """Korean field with English fallback: pick(p,'title',ko) → title_kr or title."""
    if ko and d.get(f"{key}_kr"):
        return d[f"{key}_kr"]
    return d.get(key, default)


# ── date helpers ──────────────────────────────────────────────────────
MONTHS = ["", "Jan.", "Feb.", "Mar.", "Apr.", "May", "Jun.",
          "Jul.", "Aug.", "Sep.", "Oct.", "Nov.", "Dec."]


def ym(v, ko=False) -> str:
    """2020-04 → 'Apr. 2020' (en) / '2020.04' (ko).  present → Present / 현재."""
    s = str(v)
    if s.lower() in ("present", "current", ""):
        return "현재" if ko else "Present"
    parts = s.split("-")
    if len(parts) < 2:
        return s
    return f"{parts[0]}.{parts[1]}" if ko else f"{MONTHS[int(parts[1])]} {parts[0]}"


def dot(v, ko=False) -> str:
    """Numeric period style used in the Research Projects list."""
    s = str(v)
    if s.lower() in ("present", "current", ""):
        return "현재" if ko else "Present"
    parts = s.split("-")
    return f"{parts[0]}.{parts[1]}" if len(parts) >= 2 else s


def span(a, b, ko=False) -> str:
    return f"{ym(a, ko)} – {ym(b, ko)}" if not ko else f"{ym(a, ko)} ~ {ym(b, ko)}"


# ── paragraph helpers ─────────────────────────────────────────────────
def _style_run(r, size, bold, italic):
    r.font.size = Pt(size)
    r.bold = bold
    r.italic = italic
    r._element.get_or_add_rPr().get_or_add_rFonts().set(qn("w:eastAsia"), KO_FONT)


def para(doc_or_cell, text="", *, size=BODY_PT, bold=False, italic=False,
         space_after=2, indent=0.0, hanging=0.0, align=None):
    p = doc_or_cell.add_paragraph()
    pf = p.paragraph_format
    pf.space_before = Pt(0)
    pf.space_after = Pt(space_after)
    pf.line_spacing = 1.04
    if indent:
        pf.left_indent = Pt(indent)
    if hanging:
        pf.first_line_indent = Pt(-hanging)
    p.alignment = align if align is not None else WD_ALIGN_PARAGRAPH.LEFT
    if text:
        _style_run(p.add_run(text), size, bold, italic)
    return p


def rich(doc, chunks, *, size=BODY_PT, space_after=2, indent=18, hanging=18):
    """chunks: list of (text, {'bold':bool,'italic':bool})."""
    p = para(doc, "", size=size, space_after=space_after, indent=indent, hanging=hanging)
    for text, fmt in chunks:
        _style_run(p.add_run(text), size, fmt.get("bold", False), fmt.get("italic", False))
    return p


# Width reserved for the "period:" prefix in date-led lists, so a wrapped line
# resumes where the text begins rather than under the date.
TAB_PROJECT = 96
TAB_HISTORY = {"en": 112, "ko": 96}


def dated(doc, prefix, chunks, width, *, size=BODY_PT, space_after=2):
    """'2026.01 ~ 2030.12:' <tab> body, with the body hanging-indented to `width`."""
    p = doc.add_paragraph()
    pf = p.paragraph_format
    pf.space_before = Pt(0)
    pf.space_after = Pt(space_after)
    pf.line_spacing = 1.04
    pf.left_indent = Pt(width)
    pf.first_line_indent = Pt(-width)
    pf.tab_stops.add_tab_stop(Pt(width))
    p.alignment = WD_ALIGN_PARAGRAPH.LEFT
    _style_run(p.add_run(f"{prefix}\t"), size, False, False)
    for text, fmt in chunks:
        _style_run(p.add_run(text), size, fmt.get("bold", False), fmt.get("italic", False))
    return p


def section(doc, title):
    p = para(doc, f"■ {title}", size=HEAD_PT, bold=True, space_after=3)
    p.paragraph_format.space_before = Pt(9)
    return p


def _photo_stream(path: Path, px_h: int = 420):
    """Downscale the ID photo to ~300 dpi for the 1.09 x 1.4 in slot.

    The source 증명사진 is 1570x2019; embedding it raw makes each .docx ~1.3 MB.
    Falls back to the original bytes if Pillow is unavailable.
    """
    try:
        from PIL import Image
    except ImportError:
        return str(path)
    im = Image.open(path).convert("RGB")
    im = im.resize((round(px_h * im.width / im.height), px_h), Image.LANCZOS)
    buf = io.BytesIO()
    im.save(buf, format="PNG", optimize=True)
    buf.seek(0)
    return buf


# ── header (text left, ID photo right) ────────────────────────────────
def build_header(doc, profile, ko, web=False):
    photo = ROOT / profile.get("photo", "")
    # The web edition drops the ID photo — the published site and the Typst CV
    # carry no portrait, and source/ is never pushed anyway.
    has_photo = photo.is_file() and not web

    tbl = doc.add_table(rows=1, cols=2 if has_photo else 1)
    tbl.alignment = WD_TABLE_ALIGNMENT.LEFT
    tbl.autofit = False
    # Fixed layout, or Word/LibreOffice re-flows the columns and wraps the
    # contact block into the photo's gutter.
    tbl._tbl.tblPr.append(OxmlElement("w:tblLayout", {qn("w:type"): "fixed"}))
    # Widths must land on the tblGrid too (table.columns), not just the cells —
    # otherwise LibreOffice/Word keeps the equal split and wraps the contact block.
    if has_photo:
        tbl.columns[0].width = Inches(5.45)
        tbl.columns[1].width = Inches(1.35)
    else:
        tbl.columns[0].width = Inches(6.8)
    left = tbl.cell(0, 0)
    left.paragraphs[0]._element.getparent().remove(left.paragraphs[0]._element)

    name = (f"{profile['name_kr']} ({profile['name']}), 공학박사" if ko
            else f"{profile['name']} ({profile['name_kr']}), {profile['honorific']}")
    p = para(left, "", space_after=3)
    _style_run(p.add_run(name), NAME_PT, True, False)

    c = profile["contact"]
    if ko:
        lines = [
            f"{profile['affiliation_kr']} {profile['division_kr']}",
            f"직위: {profile['title_kr']}",
            f"사무실: {c['office_phone']}   휴대전화: {c['mobile']}",
            f"이메일: {c['email']}   국적: {profile['nationality_kr']}",
            f"웹: {profile['links']['website']}",
        ]
    else:
        lines = [
            f"{profile['affiliation']}, {profile['division']}",
            f"Position: {profile['title']}",
            f"Office: {c['office_phone']}   Mobile: {c['mobile']}",
            f"E-mail: {c['email']}   Nationality: {profile['nationality']}",
            f"Web: {profile['links']['website']}",
            f"Scholar: {profile['links']['scholar']}",
        ]
    for ln in lines:
        q = para(left, ln, size=9, space_after=0)
        for r in q.runs:
            r.font.color.rgb = RGBColor(0x33, 0x33, 0x33)

    if has_photo:
        right = tbl.cell(0, 1)
        rp = right.paragraphs[0]
        rp.alignment = WD_ALIGN_PARAGRAPH.RIGHT
        rp.paragraph_format.space_after = Pt(0)
        rp.add_run().add_picture(_photo_stream(photo), width=Inches(1.09), height=Inches(1.4))

    para(doc, "", space_after=0, size=2)


# ── sections ──────────────────────────────────────────────────────────
def build_education(doc, profile, ko):
    section(doc, "학력" if ko else "Education")
    w = TAB_HISTORY["ko" if ko else "en"]
    for ed in profile["education"]:
        if ko:
            body = (f"{pick(ed,'institution',ko)} {pick(ed,'field',ko)} "
                    f"{pick(ed,'degree',ko)}")
            if ed.get("advisor_kr"):
                body += f" (지도교수: {ed['advisor_kr']})"
        else:
            body = f"{ed['degree']} {ed['field']}, {ed['institution']}"
            if ed.get("advisor"):
                body += f" (Advisor: {ed['advisor']})"
        dated(doc, f"{span(ed['start'], ed['end'], ko)}:", [(body, {})], w)


def build_experience(doc, profile, ko):
    section(doc, "경력" if ko else "Experience")
    w = TAB_HISTORY["ko" if ko else "en"]
    for ex in profile["experience"]:
        if ko:
            org = pick(ex, "org", ko)
            if ex.get("division_kr"):
                org += f" {ex['division_kr']}"
            body = f"{org}, {pick(ex,'role',ko)}"
            if ex.get("pi_kr"):
                body += f" (지도: {ex['pi_kr']})"
            if ex.get("note_kr"):
                body += f" — {ex['note_kr']}"
        else:
            body = f"{ex['role']}, {ex['org']}"
            if ex.get("division"):
                body += f", {ex['division']}"
            if ex.get("pi"):
                body += f" (PI: {ex['pi']})"
            if ex.get("note"):
                body += f" — {ex['note']}"
        dated(doc, f"{span(ex['start'], ex['end'], ko)}:", [(body, {})], w)


def pub_chunks(n, p):
    authors = ", ".join(p.get("authors", []))
    bib = [str(p[k]) for k in ("volume", "issue", "pages") if p.get(k)]
    out = [(f"[{n}] {authors}, ", {}),
           (f"“{p['title']}”, ", {}),
           (p["venue"], {"italic": True})]
    if bib:
        out.append((", " + ", ".join(bib), {}))
    out.append((f" ({p['year']})", {}))
    return out


def build_publications(doc, pubs):
    first = sorted([p for p in pubs if p.get("role") in ("first", "co_first", "corresponding")],
                   key=lambda p: -(p.get("year") or 0))
    co = sorted([p for p in pubs if p.get("role") == "co"], key=lambda p: -(p.get("year") or 0))

    section(doc, "First Author Peer-Reviewed Publications "
                 "(+: equal contributions, *: corresponding author)")
    for i, p in enumerate(first, 1):
        rich(doc, pub_chunks(i, p))

    section(doc, "Co-Author Peer-Reviewed Publications")
    for i, p in enumerate(co, 1):
        rich(doc, pub_chunks(i, p))


def build_projects(doc, projects, ko):
    section(doc, "연구과제" if ko else "Research Projects")
    for p in sorted(projects, key=lambda x: str(x.get("start", "")), reverse=True):
        prefix = f"{dot(p['start'], ko)} ~ {dot(p['end'], ko)}:"
        q = ("「", "」") if ko else ("“", "”")
        chunks = [(f"{q[0]}{pick(p,'title',ko)}{q[1]}", {}),
                  (f", {pick(p,'sponsor',ko)} ({pick(p,'role',ko)})", {})]
        dated(doc, prefix, chunks, TAB_PROJECT)


ROLE_LABEL = {"oral": "Oral Presenter", "poster": "Poster",
              "virtual_poster": "Virtual Poster", "invited": "Invited Talk",
              "keynote": "Keynote"}


def build_talks(doc, talks):
    for title, kind in (("International Conference", "international"),
                        ("Domestic Conference", "domestic")):
        section(doc, title)
        rows = sorted([t for t in talks if t.get("type") == kind],
                      key=lambda t: str(t.get("date", "")), reverse=True)
        for i, t in enumerate(rows, 1):
            rich(doc, [(f"[{i}] {', '.join(t.get('authors', []))}, ", {}),
                       (f"“{t['title']}”, ", {}),
                       (t["venue"], {"italic": True}),
                       (f", {t.get('location', '')}, {ym(t['date'])} "
                        f"({ROLE_LABEL.get(t.get('role', ''), t.get('role', ''))})", {})])


def build_patents(doc, patents, ko):
    section(doc, "특허" if ko else "Patents")
    for i, p in enumerate(sorted(patents, key=lambda x: -(x.get("year") or 0)), 1):
        inv = ", ".join(p.get("inventors", []))
        title = p["title"] if ko else (p.get("title_en") or p["title"])
        tail = f"{p['number']}, {pick(p,'jurisdiction',ko)}"
        if p.get("assignee"):
            label = "출원인" if ko else "assignee"
            tail += f", {label}: {pick(p,'assignee',ko)}"
        quote = ("「", "」") if ko else ("“", "”")
        rich(doc, [(f"[{i}] {inv}, ", {}),
                   (f"{quote[0]}{title}{quote[1]}, ", {}),
                   (f"{tail} ({p['year']})", {})])


SPEC_LABELS = {
    "forward_simulation": ("Forward Simulation", "나노포토닉스 정방향 시뮬레이션"),
    "learning_and_optimization": ("Learning & Optimization", "딥러닝 및 최적화"),
    "manufacturing_apps": ("Manufacturing Applications", "제조 응용"),
    "nanophotonics_apps": ("Nanophotonics Applications", "나노포토닉스 응용"),
    "autonomous_lab": ("Metal Autonomous Laboratory", "금속 소재 자율실험실"),
}


def build_specializations(doc, profile, ko):
    section(doc, "전문 분야" if ko else "Specializations")
    spec = profile["specializations_kr" if ko else "specializations"]
    for key, (en, kr) in SPEC_LABELS.items():
        if key not in spec:
            continue
        rich(doc, [(f"{kr if ko else en}. ", {"bold": True}),
                   (", ".join(spec[key]), {})], indent=18, hanging=0)


def build_collaborations(doc, profile, ko):
    section(doc, "공동연구 그룹" if ko else "Active Collaborations")
    for c in profile["collaborations"]:
        rich(doc, [(f"{pick(c,'lab',ko)}", {"bold": True}),
                   (f", {pick(c,'institution',ko)} — "
                    f"{'PI' if not ko else '책임자'}: {pick(c,'pi',ko)}. "
                    f"{pick(c,'topic',ko)}", {})], indent=18, hanging=18)


def build_interests(doc, profile, ko):
    section(doc, "연구 관심분야" if ko else "Research Interests")
    items = profile["research_interests_kr" if ko else "research_interests"]
    para(doc, ", ".join(items) + ("" if ko else "."), indent=18)


def build_scholar(doc, profile):
    m = profile.get("scholar_metrics") or {}
    if not m:
        return
    section(doc, "Google Scholar Metrics")
    rich(doc, [("Total citations: ", {}), (str(m['total_citations']), {"bold": True}),
               (" · h-index: ", {}), (str(m['h_index']), {"bold": True}),
               (" · i10-index: ", {}), (str(m['i10_index']), {"bold": True}),
               (f" · as of {m['last_updated']}.", {})], indent=18, hanging=0)


# ── document assembly ─────────────────────────────────────────────────
def build(lang: str, out: Path, web: bool = False):
    ko = lang == "ko"
    profile = load("profile.yaml")

    doc = Document()
    normal = doc.styles["Normal"]
    normal.font.size = Pt(BODY_PT)
    normal.element.get_or_add_rPr().get_or_add_rFonts().set(qn("w:eastAsia"), KO_FONT)
    for sec in doc.sections:
        sec.top_margin = sec.bottom_margin = Pt(48)
        sec.left_margin = sec.right_margin = Pt(50)

    build_header(doc, profile, ko, web)
    build_education(doc, profile, ko)
    build_experience(doc, profile, ko)
    if not ko:
        build_publications(doc, load("publications.yaml")["publications"])
    build_projects(doc, load("projects.yaml")["projects"], ko)
    if not ko:
        build_talks(doc, load("talks.yaml")["talks"])
    build_patents(doc, load("patents.yaml")["patents"], ko)
    build_specializations(doc, profile, ko)
    build_collaborations(doc, profile, ko)
    build_interests(doc, profile, ko)
    if not ko:
        build_scholar(doc, profile)

    doc.save(out)
    print(f"OK  {out.name}  ({out.stat().st_size:,} bytes)")


def main():
    ap = argparse.ArgumentParser()
    ap.add_argument("--lang", choices=["en", "ko", "both"], default="both")
    ap.add_argument("--out", default=None, help="explicit path (single language only)")
    ap.add_argument("--web", action="store_true",
                    help="public edition: omit the ID photo")
    args = ap.parse_args()

    if args.out:
        if args.lang == "both":
            ap.error("--out requires --lang en or --lang ko")
        build(args.lang, Path(args.out), args.web)
        return

    targets = ["en", "ko"] if args.lang == "both" else [args.lang]
    for lang in targets:
        name = (f"Curriculum Vitae Chihun Lee ({STAMP}).docx" if lang == "en"
                else f"이력서 이치헌 ({STAMP}).docx")
        build(lang, ROOT / "CV" / name, args.web)


if __name__ == "__main__":
    main()
